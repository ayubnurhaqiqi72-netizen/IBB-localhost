import streamlit as st
import subprocess
import os
import json
import time
import textwrap
from datetime import datetime
import openpyxl
import pandas as pd
from pypdf import PdfReader
import docx
from google import genai
from google.genai import types

# Configuration & Rclone Config Generator

RCLONE_CONF_PATH = os.path.expanduser("~/.config/rclone/rclone.conf")

# Otomatis buat file rclone.conf di server Cloud jika ada di st.secrets
if "RCLONE_CONF" in st.secrets:
    os.makedirs(os.path.dirname(RCLONE_CONF_PATH), exist_ok=True)
    with open(RCLONE_CONF_PATH, "w", encoding="utf-8") as f:
        f.write(str(st.secrets["RCLONE_CONF"]).strip())

# Helper function untuk mengeksekusi perintah rclone
def run_rclone_command(cmd_list):
    full_cmd = [cmd_list[0], "--config", RCLONE_CONF_PATH] + cmd_list[1:]
    result = subprocess.run(full_cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise Exception(f"Rclone Error (Exit {result.returncode}): {result.stderr.strip()}")
    return result.stdout

# 1. Page Configuration & Custom CSS

st.set_page_config(
    page_title="Innovation Idea Bank",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
    /* Global White Background */
    .stApp {
        background-color: #FFFFFF;
        color: #404040;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }

    /* Top Banner / Header Card */
    .header-card {
        background-color: #FFFFFF;
        padding: 12px 24px;
        border-radius: 4px;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.08);
        border-bottom: 3px solid #EA0029;
        margin-bottom: 20px;
        display: flex;
        align-items: center;
        justify-content: space-between;
    }

    .title {
        font-size: 1.8rem;
        font-weight: 700;
        color: #404040 !important;
        margin: 0;
    }

    .subtitle {
        font-size: 0.85rem;
        color: #6D6E71 !important;
        margin-top: 4px;
    }

    /* Dark Sidebar Styling */
    section[data-testid="stSidebar"] {
        background-color: #262626 !important;
        border-right: 1px solid #404040;
    }
    section[data-testid="stSidebar"] * {
        color: #FFFFFF !important;
    }
    div[data-testid="stSidebarNav"] {
        background-color: transparent;
    }

    /* Radio Button Styling in Sidebar */
    div[role="radiogroup"] > label {
        padding: 6px 0px !important;
        background-color: transparent !important;
        border: none !important;
        box-shadow: none !important;
    }
    div[role="radiogroup"] > label:hover {
        background-color: transparent !important;
        color: #EA0029 !important;
    }

    /* Card Container */
    .content-card {
        background-color: #FFFFFF;
        border: 1px solid #E5E7EB;
        border-radius: 6px;
        padding: 20px;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.05);
        margin-bottom: 20px;
    }

    /* Red Primary Action Buttons */
    .stButton > button {
        background-color: #EA0029 !important;
        color: #FFFFFF !important;
        border-radius: 4px !important;
        padding: 0.5rem 1.5rem !important;
        font-weight: 600 !important;
        border: none !important;
        box-shadow: 0 2px 4px rgba(234, 0, 41, 0.2);
        transition: all 0.2s;
    }
    .stButton > button:hover {
        background-color: #C40022 !important;
        box-shadow: 0 4px 8px rgba(234, 0, 41, 0.3);
    }
    .stButton > button:disabled {
        background-color: #C40022 !important;
        color: #FFFFFF !important;
        opacity: 0.85 !important;
        cursor: not-allowed !important;
    }

    /* Dividers */
    hr {
        border-color: #EA0029;
        opacity: 0.3;
        margin: 15px 0;
    }

    /* Teks Label File Uploader */
    div[data-testid="stFileUploader"] label p {
        color: #404040 !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
    }

    /* Teks Label Text Input */
    div[data-testid="stTextInput"] label p {
        color: #404040 !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
    }

    /* Teks dalam Kotak Info Alert */
    div[data-testid="stAlert"] p {
        color: #6D6E71 !important;
        font-weight: 500 !important;
    }

    /* Progress Bar Blue (Overall Progress) */
    div.blue-bar div[data-testid="stProgress"] > div > div {
        background-color: #0070F3 !important;
        height: 8px !important;
        border-radius: 4px !important;
    }
    div.blue-bar div[data-testid="stProgress"] > div {
        background-color: #E2E8F0 !important;
        height: 8px !important;
        border-radius: 4px !important;
    }

    /* Progress Bar Green (Per-file Progress) */
    div.green-bar div[data-testid="stProgress"] > div > div {
        background-color: #10B981 !important;
        height: 8px !important;
        border-radius: 4px !important;
    }
    div.green-bar div[data-testid="stProgress"] > div {
        background-color: #E2E8F0 !important;
        height: 8px !important;
        border-radius: 4px !important;
    }

    /* Custom Minimalist Notification Containers */
    .status-box-success {
        background-color: #F0FDF4;
        border-left: 3px solid #16A34A;
        border-top: 1px solid #DCFCE7;
        border-right: 1px solid #DCFCE7;
        border-bottom: 1px solid #DCFCE7;
        padding: 8px 12px;
        border-radius: 4px;
        font-size: 0.82rem;
        color: #404040;
        margin-top: 8px;
        margin-bottom: 10px;
    }

    .status-box-error {
        background-color: #FEF2F2;
        border-left: 3px solid #EA0029;
        border-top: 1px solid #FEE2E2;
        border-right: 1px solid #FEE2E2;
        border-bottom: 1px solid #FEE2E2;
        padding: 8px 12px;
        border-radius: 4px;
        font-size: 0.82rem;
        color: #404040;
        margin-top: 8px;
        margin-bottom: 10px;
    }
</style>
""", unsafe_allow_html=True)

# API Key & Path Configurations
GEMINI_API_KEY = st.secrets.get("GEMINI_API_KEY", "")
RCLONE_EXE_PATH = "rclone"
RCLONE_REMOTE_NAME = "innovation_idea_bank:"
EXCEL_FILENAME = "20260813 Innovation_Idea_Bank_Source(AutoRecovered).xlsx"

# Initialize Gemini Client
client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

# Initialize Session State
if "upload_history" not in st.session_state:
    st.session_state["upload_history"] = []
if "is_processing" not in st.session_state:
    st.session_state["is_processing"] = False


# 2. Unified System Instruction for AI Engine

UNIFIED_SYSTEM_INSTRUCTION = """
You are an expert technology innovation proposal reviewer and data processing assistant.
Your task is to thoroughly read the provided PDF proposal/essay document text and extract structured information.

RULES FOR EXTRACTION:
1. Treat this document independently without inventing facts. Write "Not mentioned in document." if missing.
2. Tag logical inferences with "[Inference]".
3. Retain technical terms and quantitative figures as written.
4. For text block values in 'raw_staging': EACH VALUE MUST BE A SINGLE COHESIVE PARAGRAPH (DO NOT USE BULLET POINTS OR NUMBERING).

Return ONLY JSON format with two root keys: "idea_master" and "raw_staging".

Expected JSON Structure:
{
  "idea_master": {
    "Year": 2026,
    "Institution_Company": "...",
    "Business_Area": "...",
    "Innovation_Type": "...",
    "Technology": "...",
    "Benefit_Category": "..."
  },
  "raw_staging": {
    "Source_Type": "...",
    "Source_Name": "...",
    "Source_Category": "...",
    "Team_Participant": "...",
    "Idea_Title_as_Written": "...",
    "Problem_as_Written": "Cohesive paragraph...",
    "Solution_as_Written": "Cohesive paragraph...",
    "Technology_Mentioned": "Cohesive paragraph...",
    "Benefit_Mentioned": "Cohesive paragraph...",
    "Stage_Prototype_Info": "Cohesive paragraph...",
    "Data_Needed": "Cohesive paragraph...",
    "Key_Metrics": "Cohesive paragraph..."
  }
}
"""


# 3. Helper Functions

def extract_text_from_file(file_path):
    file_ext = os.path.splitext(file_path)[1].lower()
    text = ""
    if file_ext == ".pdf":
        reader = PdfReader(file_path)
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    elif file_ext in [".docx", ".doc"]:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text


def analyze_document_content(text_content):
    max_retries = 5
    base_delay = 5

    for attempt in range(1, max_retries + 1):
        try:
            response = client.models.generate_content(
                model='gemini-3.6-flash',
                contents=f"Extract structured data from the following document text:\n\n{text_content}",
                config=types.GenerateContentConfig(
                    system_instruction=UNIFIED_SYSTEM_INSTRUCTION,
                    response_mime_type="application/json",
                    temperature=0.1
                ),
            )
            return json.loads(response.text)
        except Exception as e:
            err_str = str(e)
            if ("503" in err_str or "429" in err_str or "UNAVAILABLE" in err_str) and attempt < max_retries:
                sleep_time = base_delay * (2 ** (attempt - 1))
                time.sleep(sleep_time)
            else:
                raise e


def find_next_available_row(sheet, col_idx=2, start_row=4):
    for r in range(start_row, 2000):
        val = sheet.cell(row=r, column=col_idx).value
        if val is None or str(val).strip() == "" or str(val).startswith("="):
            return r
    return sheet.max_row + 1


def find_master_row_by_raw_id(sheet_master, target_raw_id):
    for r in range(4, sheet_master.max_row + 1):
        cell_val = sheet_master.cell(row=r, column=2).value
        if cell_val and str(cell_val).strip() == str(target_raw_id).strip():
            return r
        if target_raw_id and "RAW-2026-" in str(target_raw_id):
            try:
                raw_num = int(str(target_raw_id).split("-")[-1])
                expected_row = raw_num + 3
                if r == expected_row:
                    return r
            except ValueError:
                pass
    return None


def process_excel_database_update(excel_path, doc_filename, doc_text):
    wb = openpyxl.load_workbook(excel_path)
    sheet_staging = wb["Raw_Staging"]
    sheet_master = wb["Idea_Master"]

    existing_staging_row = None
    raw_id = None

    for row in range(4, sheet_staging.max_row + 1):
        cell_file = sheet_staging.cell(row=row, column=2).value
        if cell_file and str(cell_file).strip().lower() == doc_filename.strip().lower():
            existing_staging_row = row
            raw_id = sheet_staging.cell(row=row, column=1).value
            break

    if not existing_staging_row:
        extracted_json = analyze_document_content(doc_text)
        data_a = extracted_json.get("idea_master", {})
        data_b = extracted_json.get("raw_staging", {})

        new_staging_row = find_next_available_row(
            sheet_staging,
            col_idx=2,
            start_row=4
        )
        existing_raw_id = sheet_staging.cell(
            row=new_staging_row,
            column=1
        ).value

        if existing_raw_id and str(existing_raw_id).strip() != "" and not str(existing_raw_id).startswith("="):
            raw_id = str(existing_raw_id).strip()
        else:
            raw_id = f"RAW-2026-{new_staging_row - 3:03d}"
            sheet_staging.cell(
                row=new_staging_row,
                column=1,
                value=raw_id
            )

        sheet_staging.cell(row=new_staging_row, column=2, value=doc_filename)
        sheet_staging.cell(row=new_staging_row, column=3, value=data_b.get("Source_Type"))
        sheet_staging.cell(row=new_staging_row, column=4, value=data_b.get("Source_Name"))
        sheet_staging.cell(row=new_staging_row, column=5, value=data_b.get("Source_Category"))
        sheet_staging.cell(row=new_staging_row, column=6, value=data_b.get("Team_Participant"))
        sheet_staging.cell(row=new_staging_row, column=7, value=data_b.get("Idea_Title_as_Written"))
        sheet_staging.cell(row=new_staging_row, column=8, value=data_b.get("Problem_as_Written"))
        sheet_staging.cell(row=new_staging_row, column=9, value=data_b.get("Solution_as_Written"))
        sheet_staging.cell(row=new_staging_row, column=10, value=data_b.get("Technology_Mentioned"))
        sheet_staging.cell(row=new_staging_row, column=11, value=data_b.get("Benefit_Mentioned"))
        sheet_staging.cell(row=new_staging_row, column=12, value=data_b.get("Stage_Prototype_Info"))
        sheet_staging.cell(row=new_staging_row, column=13, value=data_b.get("Data_Needed"))
        sheet_staging.cell(row=new_staging_row, column=14, value=data_b.get("Key_Metrics"))

        master_target_row = find_master_row_by_raw_id(sheet_master, raw_id)

        target_row = (
            master_target_row
            if master_target_row
            else find_next_available_row(
                sheet_master,
                col_idx=2,
                start_row=4
            )
        )

        existing_idea_id = sheet_master.cell(row=target_row, column=1).value

        if not existing_idea_id or str(existing_idea_id).strip() == "" or str(existing_idea_id).startswith("="):
            new_idea_id = f"IDEA-2026-{target_row - 3:04d}"
            sheet_master.cell(
                row=target_row,
                column=1,
                value=new_idea_id
            )

        sheet_master.cell(row=target_row, column=2, value=raw_id)
        sheet_master.cell(row=target_row, column=3, value=data_b.get("Idea_Title_as_Written"))
        sheet_master.cell(row=target_row, column=4, value=data_b.get("Source_Type"))
        sheet_master.cell(row=target_row, column=5, value=data_b.get("Source_Name"))
        sheet_master.cell(row=target_row, column=6, value=data_b.get("Source_Category"))
        sheet_master.cell(row=target_row, column=7, value=data_a.get("Year"))
        sheet_master.cell(row=target_row, column=8, value=data_b.get("Team_Participant"))
        sheet_master.cell(row=target_row, column=9, value=data_a.get("Institution_Company"))
        sheet_master.cell(row=target_row, column=12, value=data_a.get("Business_Area"))
        sheet_master.cell(row=target_row, column=13, value=data_a.get("Innovation_Type"))
        sheet_master.cell(row=target_row, column=14, value=data_a.get("Technology"))
        sheet_master.cell(row=target_row, column=15, value=data_a.get("Benefit_Category"))

    wb.save(excel_path)
    wb.close()


# 4. Sidebar Menu

st.sidebar.markdown(
    "<h2 style='font-size: 1.2rem; font-weight: 700; color: #FFFFFF;'>MENU</h2>",
    unsafe_allow_html=True
)

menu = st.sidebar.radio(
    "Select Module:",
    ["Upload & Process", "Search & View Database"],
    label_visibility="collapsed"
)


# 5. Top Header Banner

st.markdown("""<div class="header-card">
<div>
<div class="title">INNOVATION IDEA BANK</div>
<div class="subtitle">Automated Ingestion & Cloud Database Integration System</div>
</div>
</div>""", unsafe_allow_html=True)


# 6. Menu 1: Upload & Process

if menu == "Upload & Process":
    st.markdown("""<div class="content-card">
<h3 style="color: #404040; margin-top: 0; font-size: 1.1rem; font-weight: 600;">Document Batch Ingestion</h3>
<p style="color: #6D6E71; font-size: 0.9rem;">Select one or multiple proposal/essay files (PDF or DOCX) to extract information and sync directly to the Cloud database.</p>
</div>""", unsafe_allow_html=True)

    uploaded_files = st.file_uploader(
        "Upload Proposal or Essay Documents (PDF, DOCX)",
        type=["pdf", "docx"],
        accept_multiple_files=True
    )

    if uploaded_files:

        st.markdown(
            f"""<div style="background-color: #F8FAFC; border-left: 4px solid #EA0029; padding: 10px 15px; border-radius: 4px; margin-bottom: 15px;">
<span style="color: #404040; font-weight: 600;">Total documents selected: {len(uploaded_files)} file(s)</span>
</div>""",
            unsafe_allow_html=True
        )

        btn_label = "Processing....." if st.session_state["is_processing"] else "Start Processing"
        
        if st.button(btn_label, disabled=st.session_state["is_processing"]):
            st.session_state["is_processing"] = True
            
            temp_dir = "temp_processing"
            os.makedirs(temp_dir, exist_ok=True)
            local_excel_path = os.path.join(temp_dir, EXCEL_FILENAME)
            
            st.markdown(
                """<div style="font-size: 0.82rem; color: #6D6E71; margin-bottom: 8px;">
                Downloading latest Cloud Database...
                </div>""",
                unsafe_allow_html=True
            )
            
            cmd_download_excel = [
                RCLONE_EXE_PATH,
                "copyto",
                f"{RCLONE_REMOTE_NAME}{EXCEL_FILENAME}",
                local_excel_path
            ]

            try:
                run_rclone_command(cmd_download_excel)

            except Exception as e:
                st.markdown(
                    f"""<div class="status-box-error">
                    <b>Fetch Failed:</b> Could not download database from Cloud Storage. Details: {e}
                    </div>""",
                    unsafe_allow_html=True
                )
                st.session_state["is_processing"] = False
                st.stop()

            # Dynamic Progress Container (Double Progress Bars)
            progress_container = st.container()
            with progress_container:
                file_title_placeholder = st.empty()
                
                st.markdown('<div class="blue-bar">', unsafe_allow_html=True)
                overall_progress = st.progress(0)
                st.markdown('</div>', unsafe_allow_html=True)
                
                st.markdown('<div class="green-bar">', unsafe_allow_html=True)
                file_progress = st.progress(0)
                st.markdown('</div>', unsafe_allow_html=True)
                
                status_box_placeholder = st.empty()

            total_files = len(uploaded_files)

            for index, file_item in enumerate(uploaded_files, start=1):
                file_title_placeholder.markdown(
                    f"<h5 style='color: #404040; font-size: 0.95rem; font-weight: 600; margin-top: 15px; margin-bottom: 8px;'>Processing File [{index}/{total_files}]: {file_item.name}</h5>",
                    unsafe_allow_html=True
                )

                local_doc_path = os.path.join(temp_dir, file_item.name)
                current_time_str = datetime.now().strftime("%H:%M\n%d/%m/%Y")

                try:
                    file_progress.progress(20)

                    with open(local_doc_path, "wb") as f:
                        f.write(file_item.getbuffer())

                    file_progress.progress(40)

                    cmd_upload_doc = [
                        RCLONE_EXE_PATH,
                        "copy",
                        local_doc_path,
                        f"{RCLONE_REMOTE_NAME}Raw_Documents"
                    ]
                    run_rclone_command(cmd_upload_doc)

                    file_progress.progress(60)
                    doc_text = extract_text_from_file(local_doc_path)

                    if doc_text.strip():
                        file_progress.progress(80)
                        process_excel_database_update(
                            local_excel_path,
                            file_item.name,
                            doc_text
                        )
                        file_progress.progress(100)
                        
                        status_box_placeholder.markdown(
                            f"""<div class="status-box-success">
                            <b>Success:</b> File processed and mapped successfully.
                            </div>""",
                            unsafe_allow_html=True
                        )

                        st.session_state["upload_history"].append({
                            "File Name": file_item.name,
                            "Date": current_time_str,
                            "Status": "Success"
                        })
                    else:
                        file_progress.progress(100)
                        err_msg = "Warning: Text Extraction Empty (File encrypted or scan-only)"
                        status_box_placeholder.markdown(
                            f"""<div class="status-box-error">
                            <b>Warning:</b> Could not extract text from file <b>{file_item.name}</b>. File may be encrypted or scan-only.
                            </div>""",
                            unsafe_allow_html=True
                        )
                        st.session_state["upload_history"].append({
                            "File Name": file_item.name,
                            "Date": current_time_str,
                            "Status": f"Error ({err_msg})"
                        })

                except Exception as e:
                    file_progress.progress(100)
                    err_type = type(e).__name__
                    status_box_placeholder.markdown(
                        f"""<div class="status-box-error">
                        <b>Processing Error:</b> Failed to process <b>{file_item.name}</b>. Details: {e}
                        </div>""",
                        unsafe_allow_html=True
                    )
                    st.session_state["upload_history"].append({
                        "File Name": file_item.name,
                        "Date": current_time_str,
                        "Status": f"Error ({err_type}: {e})"
                    })

                finally:
                    if os.path.exists(local_doc_path):
                        os.remove(local_doc_path)
                    
                    overall_progress.progress(int((index / total_files) * 100))

            st.markdown(
                """<div style="font-size: 0.82rem; color: #6D6E71; margin-top: 15px; margin-bottom: 8px;">
                Syncing updated database back to Cloud Repository...
                </div>""",
                unsafe_allow_html=True
            )
            
            cmd_upload_excel = [
                RCLONE_EXE_PATH,
                "copyto",
                local_excel_path,
                f"{RCLONE_REMOTE_NAME}{EXCEL_FILENAME}"
            ]
            
            try:
                run_rclone_command(cmd_upload_excel)
                st.markdown(
                    """<div class="status-box-success" style="font-size: 0.88rem; padding: 10px 14px; margin-top: 12px;">
                    <b>Batch Process Complete:</b> All files processed and Cloud database synchronized successfully.
                    </div>""",
                    unsafe_allow_html=True
                )
            except Exception as e:
                st.markdown(
                    f"""<div class="status-box-error" style="font-size: 0.88rem; padding: 10px 14px; margin-top: 12px;">
                    <b>Cloud Sync Error:</b> Processed locally, but failed to sync back to OneDrive. Please close Excel Online if it is currently open. Details: {e}
                    </div>""",
                    unsafe_allow_html=True
                )

            if os.path.exists(local_excel_path):
                os.remove(local_excel_path)
            
            st.session_state["is_processing"] = False
            st.rerun()

    # Section: History Uploaded File Table
    st.markdown("""<div class="content-card" style="margin-top: 25px;">
<h3 style="color: #404040; margin-top: 0; font-size: 1.1rem; font-weight: 600;">History Uploaded File</h3>
</div>""", unsafe_allow_html=True)

    if st.session_state["upload_history"]:
        df_history = pd.DataFrame(st.session_state["upload_history"])
        df_history.index = range(1, len(df_history) + 1)
        df_history.index.name = "No"
        
        html_table = df_history.to_html(classes="custom-history-table", escape=False)
        
        custom_table_css = """
        <style>
            table.custom-history-table {
                width: 100%;
                border-collapse: collapse;
                background-color: #FFFFFF !important;
                color: #000000 !important;
                font-size: 0.85rem;
                border: 1px solid #E5E7EB;
                border-radius: 4px;
                overflow: hidden;
            }
            table.custom-history-table th {
                background-color: #F1F5F9 !important;
                color: #000000 !important;
                font-weight: 700 !important;
                padding: 10px 12px;
                text-align: left;
                border-bottom: 2px solid #E5E7EB;
            }
            table.custom-history-table td {
                background-color: #FFFFFF !important;
                color: #000000 !important;
                padding: 10px 12px;
                border-bottom: 1px solid #E5E7EB;
                white-space: pre-line;
            }
            table.custom-history-table tr:hover td {
                background-color: #F8FAFC !important;
            }
        </style>
        """
        
        st.markdown(custom_table_css + html_table, unsafe_allow_html=True)
    else:
        st.info("No upload history recorded yet in this session.")


# 7. Menu 2: Search & View Database

elif menu == "Search & View Database":
    st.markdown("""<div class="content-card">
<h3 style="color: #404040; margin-top: 0; font-size: 1.1rem; font-weight: 600;">Database Records & Search Engine</h3>
<p style="color: #6D6E71; font-size: 0.9rem;">Search, filter, and review structured proposal records synced live from OneDrive storage.</p>
</div>""", unsafe_allow_html=True)

    if st.button("Fetch / Refresh Database"):
        with st.spinner("Fetching latest database from OneDrive..."):
            temp_dir = "temp_view"
            os.makedirs(temp_dir, exist_ok=True)
            local_excel_path = os.path.join(temp_dir, EXCEL_FILENAME)

            cmd_download = [
                RCLONE_EXE_PATH,
                "copyto",
                f"{RCLONE_REMOTE_NAME}{EXCEL_FILENAME}",
                local_excel_path
            ]

            try:
                run_rclone_command(cmd_download)
                df_staging = pd.read_excel(
                    local_excel_path,
                    sheet_name="Idea_Master",
                    skiprows=2
                )
                st.session_state["df_database"] = df_staging
                st.markdown(
                    """<div class="status-box-success" style="margin-bottom: 15px;">
                    <b>Success:</b> Latest database loaded from OneDrive.
                    </div>""",
                    unsafe_allow_html=True
                )
            except Exception as e:
                st.markdown(
                    f"""<div class="status-box-error" style="margin-bottom: 15px;">
                    <b>Fetch Failed:</b> {e}
                    </div>""",
                    unsafe_allow_html=True
                )
            finally:
                if os.path.exists(local_excel_path):
                    os.remove(local_excel_path)

    if "df_database" in st.session_state:
        df = st.session_state["df_database"]

        search_query = st.text_input(
            "🔍 Search by Document Name, Title, or Keyword:",
            ""
        )

        if search_query.strip():
            filtered_df = df[
                df.astype(str)
                .apply(
                    lambda row: row.str.contains(
                        search_query,
                        case=False,
                        na=False
                    )
                )
                .any(axis=1)
            ]
            st.write(f"Found **{len(filtered_df)}** matching record(s):")
            st.dataframe(filtered_df, use_container_width=True)
        else:
            st.write(f"Displaying all records (**{len(df)}** total):")
            st.dataframe(df, use_container_width=True)
    else:
        st.info("Click 'Fetch / Refresh Database' above to load records from OneDrive.")

